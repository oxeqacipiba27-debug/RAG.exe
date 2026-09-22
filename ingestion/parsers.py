"""
ingestion/parsers.py — Высокопроизводительный модуль ETL-парсинга и структурирования
документов MS Word (.docx) и MS Excel (.xlsx) для Enterprise RAG 2.0.

Особенности:
1. Word (.docx):
   - Последовательное извлечение параграфов и встроенных таблиц в порядке их появления в документе.
   - Сериализация строк таблиц с явными разделителями ("Cell1 | Cell2 | Cell3").
   - Нормализация пробелов и пропуск пустых блоков.
2. Excel (.xlsx):
   - Семантическое преобразование строк: "Лист: {sheet}. {Header1}: {Val1}, {Header2}: {Val2}..."
   - Чтение кэшированных расчетных значений (data_only=True) для исключения сырых формул.
   - Защита от утечек памяти (read_only=True и гарантированное закрытие книги).
   - Корректная обработка None, NaN, дат, чисел и пустых строк.
3. Чанкинг:
   - RecursiveCharacterTextSplitter для Word и текстовых данных (500-800 символов, перекрытие 50-100).
   - Атомарный или блочный чанкинг Excel без разрыва границ строк.
4. Nomic Embeddings v1.5 Protocol:
   - search_document: для индексации документов.
   - search_query: для пользовательских поисковых запросов.
"""

from __future__ import annotations

import sys
import re
import math
import hashlib
import logging
from pathlib import Path
from datetime import datetime, date
from typing import List, Dict, Any, Optional

# Настройка UTF-8 для вывода в консоль Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

logger = logging.getLogger("IngestionParsers")

# =====================================================================
# Константы протокола Nomic Embed v1.5
# =====================================================================
NOMIC_DOC_PREFIX: str = "search_document: "
NOMIC_QUERY_PREFIX: str = "search_query: "


def format_nomic_document(text: str) -> str:
    """
    Принудительно форматирует текст фрагмента документа для индексации в nomic-embed-text-v1.5.
    Префикс: строго 'search_document: ' (без дублирования).
    """
    text = text.strip()
    if text.startswith(NOMIC_DOC_PREFIX):
        return text
    if text.startswith(NOMIC_QUERY_PREFIX):
        text = text[len(NOMIC_QUERY_PREFIX):].strip()
    return f"{NOMIC_DOC_PREFIX}{text}"


def format_nomic_query(query: str) -> str:
    """
    Принудительно форматирует поисковый запрос пользователя для nomic-embed-text-v1.5.
    Префикс: строго 'search_query: ' (без дублирования).
    """
    query = query.strip()
    if query.startswith(NOMIC_QUERY_PREFIX):
        return query
    if query.startswith(NOMIC_DOC_PREFIX):
        query = query[len(NOMIC_DOC_PREFIX):].strip()
    return f"{NOMIC_QUERY_PREFIX}{query}"


# =====================================================================
# 1. Парсер MS Word (.docx)
# =====================================================================

def parse_docx(path: str) -> List[str]:
    """
    Извлекает текст из документа MS Word (.docx) последовательно.
    Обрабатывает как обычные параграфы, так и встроенные таблицы в порядке их
    следования в документе (через doc.element.body).
    
    Для строк таблиц используется разделитель ячеек ' | '.
    Пустые строки и пробельные дубликаты нормализуются.
    
    Args:
        path: Путь к файлу .docx
        
    Returns:
        Список извлеченных текстовых блоков (параграфы и строки таблиц).
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"Файл DOCX не найден: {path}")

    try:
        import docx
        from docx.text.paragraph import Paragraph
        from docx.table import Table
    except ImportError:
        err_msg = "Пакет 'python-docx' не установлен. Установите его командой: pip install python-docx"
        logger.error(err_msg)
        raise ImportError(err_msg)

    doc = docx.Document(str(file_path))
    blocks: List[str] = []

    # Обход элементов body в исходном порядке документа
    for child in doc.element.body:
        tag = child.tag.lower()
        if tag.endswith("p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            # Нормализация внутренних пробелов и табуляций
            text = re.sub(r"[ \t]+", " ", text)
            if text:
                blocks.append(text)
        elif tag.endswith("tbl"):
            table = Table(child, doc)
            for row in table.rows:
                # Извлекаем и нормализуем текст каждой ячейки
                cell_texts = []
                for cell in row.cells:
                    c_text = re.sub(r"[\r\n\t]+", " ", cell.text)
                    c_text = re.sub(r" +", " ", c_text).strip()
                    cell_texts.append(c_text)

                # Исключаем последовательные дубликаты объединенных ячеек (merged cells)
                unique_cells = []
                for c in cell_texts:
                    if not unique_cells or c != unique_cells[-1]:
                        unique_cells.append(c)

                # Собираем строку через разделитель ячеек
                row_content = [c for c in unique_cells if c]
                if row_content:
                    row_str = " | ".join(row_content)
                    blocks.append(row_str)

    logger.info(f"DOCX '{file_path.name}': извлечено {len(blocks)} структурных блоков текста.")
    return blocks


# =====================================================================
# 2. Парсер MS Excel (.xlsx)
# =====================================================================

def _format_cell_value(val: Any) -> Optional[str]:
    """Форматирует значение ячейки Excel в чистую строку без формул и NaN."""
    if val is None:
        return None

    if isinstance(val, float):
        if math.isnan(val) or math.isinf(val):
            return None
        if val.is_integer():
            return str(int(val))
        # Округление до 4 знаков без висячих нулей
        return f"{val:.4f}".rstrip("0").rstrip(".")

    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")

    val_str = str(val).strip()
    val_str = re.sub(r"[\r\n\t]+", " ", val_str)
    val_str = re.sub(r" +", " ", val_str)

    if not val_str or val_str.lower() in ("nan", "none", "null", "#n/a", "#value!", "#ref!"):
        return None

    return val_str


def parse_xlsx(path: str) -> List[str]:
    """
    Преобразует строки листов Excel (.xlsx) в семантические строки вида:
    'Лист: {sheet_name}. {Column_Header_1}: {Value_1}, {Column_Header_2}: {Value_2}...'

    Использует openpyxl в режиме:
    - data_only=True: считываются вычисленные кэшированные значения вместо формул.
    - read_only=True: оптимизированная потоковая загрузка для исключения утечек памяти.
    - Гарантированное закрытие дескриптора книги (try...finally: wb.close()).

    Args:
        path: Путь к файлу .xlsx

    Returns:
        Список семантически структурированных строк.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"Файл XLSX не найден: {path}")

    try:
        import openpyxl
    except ImportError:
        err_msg = "Пакет 'openpyxl' не установлен. Установите его командой: pip install openpyxl"
        logger.error(err_msg)
        raise ImportError(err_msg)

    semantic_rows: List[str] = []

    # Открытие с потоковым чтением и кэшированными формулами
    wb = openpyxl.load_workbook(filename=str(file_path), data_only=True, read_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            headers: Optional[List[str]] = None

            for row in ws.iter_rows(values_only=True):
                if not row or not any(v is not None for v in row):
                    continue

                # Первая непустая строка принимается за заголовки столбцов
                if headers is None:
                    raw_headers = []
                    for idx, cell in enumerate(row, 1):
                        formatted_h = _format_cell_value(cell)
                        raw_headers.append(formatted_h if formatted_h else f"Колонка_{idx}")
                    # Проверяем, что хотя бы один заголовок осмысленный
                    if any(not h.startswith("Колонка_") for h in raw_headers):
                        headers = raw_headers
                        continue
                    else:
                        headers = raw_headers

                # Преобразование строки данных в семантические пары ключ-значение
                row_items = []
                for h, cell in zip(headers, row):
                    val_str = _format_cell_value(cell)
                    if val_str is not None:
                        row_items.append(f"{h}: {val_str}")

                if row_items:
                    line = f"Лист: {sheet_name}. " + ", ".join(row_items)
                    semantic_rows.append(line)

    finally:
        wb.close()

    logger.info(f"XLSX '{file_path.name}': извлечено {len(semantic_rows)} семантических строк.")
    return semantic_rows


# =====================================================================
# 3. Унифицированный диспетчер и сплиттер (load_and_chunk_document)
# =====================================================================

def count_tokens_heuristic(text: str) -> int:
    """Оценка количества токенов (через tiktoken или fallback)."""
    try:
        import tiktoken
        enc = tiktoken.get_encoding("cl100k_base")
        return len(enc.encode(text, disallowed_special=()))
    except Exception:
        return max(1, int(len(text) / 3.5))


def load_and_chunk_document(
    file_path: str,
    chunk_size: int = 700,
    chunk_overlap: int = 75
) -> List[Dict[str, Any]]:
    """
    Унифицированный диспетчер ETL-обработки документов (.docx, .xlsx, .pdf, .txt).
    
    Для Word и текста выполняет рекурсивное разбиение (RecursiveCharacterTextSplitter)
    с сохранением смысловых границ абзацев и таблиц.
    Для Excel группирует семантические строки без разрыва границ строк.

    Returns:
        Список словарей с полями:
        - "text": чистый текст фрагмента (сырой, без префикса)
        - "metadata": словарь метаданных (source, raw_source, chunk_id, page/sheet, format, char_count, token_count)
    """
    path_obj = Path(file_path)
    if not path_obj.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")

    ext = path_obj.suffix.lower()
    source_name = path_obj.name
    doc_id = hashlib.sha256(source_name.encode("utf-8")).hexdigest()[:16]

    chunks_data: List[Dict[str, Any]] = []

    # Импорт сплиттера
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ". ", " ", ""]
        )
    except ImportError:
        # Fallback сплиттер если пакет не доступен
        class SimpleSplitter:
            def __init__(self, c_size: int, c_overlap: int):
                self.c_size = c_size
                self.c_overlap = c_overlap
            def split_text(self, text: str) -> List[str]:
                step = max(1, self.c_size - self.c_overlap)
                return [text[i : i + self.c_size].strip() for i in range(0, len(text), step) if text[i : i + self.c_size].strip()]
        splitter = SimpleSplitter(chunk_size, chunk_overlap)

    # 1. Документы MS Word (.docx)
    if ext == ".docx":
        blocks = parse_docx(str(path_obj))
        if not blocks:
            return []

        # Объединяем блоки через двойной перевод строки и делим сплиттером
        full_doc_text = "\n\n".join(blocks)
        text_chunks = splitter.split_text(full_doc_text)

        for idx, ch_text in enumerate(text_chunks, 1):
            ch_clean = ch_text.strip()
            if not ch_clean:
                continue
            meta = {
                "source": source_name,
                "raw_source": str(path_obj),
                "chunk_id": f"{doc_id}_c{idx:04d}",
                "page": 1,
                "sheet": None,
                "format": "docx",
                "char_count": len(ch_clean),
                "token_count": count_tokens_heuristic(ch_clean),
                "content_hash": hashlib.sha256(ch_clean.encode("utf-8")).hexdigest()[:16]
            }
            chunks_data.append({"text": ch_clean, "metadata": meta})

    # 2. Таблицы MS Excel (.xlsx)
    elif ext == ".xlsx":
        rows = parse_xlsx(str(path_obj))
        if not rows:
            return []

        # Группировка строк в чанки без разрыва границ строк
        current_batch: List[str] = []
        current_len = 0
        chunk_idx = 1

        def flush_batch(batch_rows: List[str], idx: int) -> Optional[Dict[str, Any]]:
            if not batch_rows:
                return None
            combined = "\n".join(batch_rows)
            # Извлекаем имя листа из первой строки батча (Лист: {name}.)
            m = re.match(r"^Лист:\s*([^.]+)\.", batch_rows[0])
            sheet_name = m.group(1).strip() if m else None

            meta = {
                "source": source_name,
                "raw_source": str(path_obj),
                "chunk_id": f"{doc_id}_c{idx:04d}",
                "page": 1,
                "sheet": sheet_name,
                "format": "xlsx",
                "char_count": len(combined),
                "token_count": count_tokens_heuristic(combined),
                "content_hash": hashlib.sha256(combined.encode("utf-8")).hexdigest()[:16]
            }
            return {"text": combined, "metadata": meta}

        for row in rows:
            row_len = len(row) + 1
            if current_len + row_len > chunk_size and current_batch:
                ch = flush_batch(current_batch, chunk_idx)
                if ch:
                    chunks_data.append(ch)
                    chunk_idx += 1
                current_batch = [row]
                current_len = len(row)
            else:
                current_batch.append(row)
                current_len += row_len

        if current_batch:
            ch = flush_batch(current_batch, chunk_idx)
            if ch:
                chunks_data.append(ch)

    # 3. Текстовые файлы (.txt)
    elif ext == ".txt":
        try:
            with open(path_obj, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read().strip()
            if content:
                text_chunks = splitter.split_text(content)
                for idx, ch_text in enumerate(text_chunks, 1):
                    ch_clean = ch_text.strip()
                    if not ch_clean:
                        continue
                    meta = {
                        "source": source_name,
                        "raw_source": str(path_obj),
                        "chunk_id": f"{doc_id}_c{idx:04d}",
                        "page": 1,
                        "sheet": None,
                        "format": "txt",
                        "char_count": len(ch_clean),
                        "token_count": count_tokens_heuristic(ch_clean),
                        "content_hash": hashlib.sha256(ch_clean.encode("utf-8")).hexdigest()[:16]
                    }
                    chunks_data.append({"text": ch_clean, "metadata": meta})
        except Exception as e:
            logger.error(f"Ошибка чтения TXT файла {source_name}: {e}")

    # 4. Документы PDF (.pdf)
    elif ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path_obj))
            chunk_idx = 1
            for p_num, page in enumerate(reader.pages, 1):
                p_text = (page.extract_text() or "").strip()
                if not p_text:
                    continue
                p_chunks = splitter.split_text(p_text)
                for ch_text in p_chunks:
                    ch_clean = ch_text.strip()
                    if not ch_clean:
                        continue
                    meta = {
                        "source": source_name,
                        "raw_source": str(path_obj),
                        "chunk_id": f"{doc_id}_c{chunk_idx:04d}",
                        "page": p_num,
                        "sheet": None,
                        "format": "pdf",
                        "char_count": len(ch_clean),
                        "token_count": count_tokens_heuristic(ch_clean),
                        "content_hash": hashlib.sha256(ch_clean.encode("utf-8")).hexdigest()[:16]
                    }
                    chunks_data.append({"text": ch_clean, "metadata": meta})
                    chunk_idx += 1
        except Exception as e:
            logger.error(f"Ошибка чтения PDF файла {source_name}: {e}")

    else:
        logger.warning(f"Неподдерживаемый формат файла для чанкинга: {ext}")

    logger.info(f"'{source_name}' ({ext}): сформировано {len(chunks_data)} чанков.")
    return chunks_data
