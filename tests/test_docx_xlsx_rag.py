"""
tests/test_docx_xlsx_rag.py — Комплексный смоук-тест пайплайна инжестии Word (.docx)
и Excel (.xlsx) с моделью nomic-ai/nomic-embed-text-v1.5.

Тестируемые компоненты:
1. Генерация синтетических файлов .docx (параграфы + таблицы) и .xlsx (формулы + мульти-листы).
2. parse_docx: последовательное извлечение текста, корректность разделителей ячеек (" | "), нормализация пробелов.
3. parse_xlsx: семантический формат ("Лист: ... {Col}: {Val}"), кэшированные значения формул, очистка None/NaN.
4. load_and_chunk_document: унифицированный диспетчер, проверка метаданных (source, chunk_id, page/sheet, format).
5. Nomic v1.5 Embedding Protocol: префиксы 'search_document: ' и 'search_query: ', идемпотентность.
6. Вычисление L2-нормализованных эмбеддингов через DenseVectorStore (trust_remote_code=True).
7. Семантический поиск и косинусное сходство (cosine similarity) на реальных запросах.
"""

from __future__ import annotations

import os
import sys
import shutil
import tempfile
import numpy as np
from pathlib import Path

# Настройка UTF-8 для вывода в консоль Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Добавляем корень проекта в sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from ingestion.parsers import (
    parse_docx,
    parse_xlsx,
    load_and_chunk_document,
    format_nomic_document,
    format_nomic_query,
    NOMIC_DOC_PREFIX,
    NOMIC_QUERY_PREFIX,
)
from rag_engine import (
    DenseVectorStore,
    DocumentIngestionPipeline,
    AppConfig,
)


def create_dummy_docx(file_path: Path):
    """Создает синтетический тестовый документ Word с параграфами и таблицей."""
    import docx

    doc = docx.Document()
    doc.add_heading("Корпоративный регламент информационной безопасности", level=1)
    doc.add_paragraph("Политика паролей: пароли пользователей должны содержать не менее 16 символов и меняться раз в 90 дней.")
    
    # Встроенная таблица
    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Роль"
    hdr_cells[1].text = "Уровень доступа"
    hdr_cells[2].text = "Ограничения"

    r1_cells = table.rows[1].cells
    r1_cells[0].text = "Администратор"
    r1_cells[1].text = "Полный доступ"
    r1_cells[2].text = "Только через аппаратный 2FA токен"

    r2_cells = table.rows[2].cells
    r2_cells[0].text = "Аналитик данных"
    r2_cells[1].text = "Только чтение"
    r2_cells[2].text = "Запрет выгрузки на внешние носители"

    doc.add_paragraph("Инциденты информационной безопасности регистрируются дежурным офицером в течение 15 минут.")
    doc.save(str(file_path))


def create_dummy_xlsx(file_path: Path):
    """Создает синтетический тестовый документ Excel с несколькими листами, формулами и пустыми ячейками."""
    import openpyxl

    wb = openpyxl.Workbook()
    
    # Лист 1: Зарплаты и сотрудники
    ws1 = wb.active
    ws1.title = "Штатное расписание"
    ws1.append(["Табельный номер", "ФИО", "Должность", "Оклад", "Премия", "Итого"])
    ws1.append([101, "Иванов Иван Иванович", "Главный архитектор AI", 350000, 50000, "=D2+E2"])
    ws1.append([102, "Петров Петр Сергеевич", "Senior DevOps Инженер", 280000, 40000, "=D3+E3"])
    ws1.append([103, "Сидорова Анна Павловна", "NLP Исследователь", 310000, None, 310000])  # None в премии

    # Лист 2: Серверная инфраструктура
    ws2 = wb.create_sheet(title="Кластер GPU")
    ws2.append(["ID Сервера", "Модель сервера", "Графические ускорители", "VRAM", "Статус"])
    ws2.append(["SRV-01", "Dell PowerEdge R750xa", "4x NVIDIA RTX 4090", "96 ГБ", "В работе"])
    ws2.append(["SRV-02", "Supermicro GPU Server", "8x NVIDIA RTX 3090", "192 ГБ", "Векторизация Nomic"])

    wb.save(str(file_path))


def cosine_similarity(a: list[float] | np.ndarray, b: list[float] | np.ndarray) -> float:
    """Вычисляет косинусное сходство двух векторов."""
    va = np.array(a, dtype=np.float32)
    vb = np.array(b, dtype=np.float32)
    na = np.linalg.norm(va)
    nb = np.linalg.norm(vb)
    if na == 0 or nb == 0:
        return 0.0
    return float(np.dot(va, vb) / (na * nb))


def run_all_tests():
    print("=" * 70)
    print("🚀 [START] Тестирование Ingestion Pipeline (DOCX, XLSX) + Nomic v1.5")
    print("=" * 70)

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        docx_file = tmp_path / "security_policy.docx"
        xlsx_file = tmp_path / "company_resources.xlsx"

        # --- 1. Генерация тестовых файлов ---
        print("\n[1/6] Генерация синтетических документов...")
        create_dummy_docx(docx_file)
        create_dummy_xlsx(xlsx_file)
        assert docx_file.exists(), "DOCX файл не создан"
        assert xlsx_file.exists(), "XLSX файл не создан"
        print(f" [OK] Файлы сгенерированы: {docx_file.name}, {xlsx_file.name}")

        # --- 2. Тестирование парсера Word (.docx) ---
        print("\n[2/6] Тестирование parse_docx...")
        docx_blocks = parse_docx(str(docx_file))
        assert len(docx_blocks) >= 4, f"Ожидалось >=4 блоков, получено {len(docx_blocks)}"
        
        # Проверка последовательности: заголовок -> параграф -> таблица -> завершающий параграф
        assert "Корпоративный регламент" in docx_blocks[0], f"Заголовок не на 1 позиции: {docx_blocks[0]}"
        assert "Политика паролей" in docx_blocks[1], f"Параграф паролей не на 2 позиции: {docx_blocks[1]}"
        
        # Проверка разделителя таблицы
        table_row_found = any(" | " in b and "Администратор" in b for b in docx_blocks)
        assert table_row_found, f"Строка таблицы с разделителем ' | ' не найдена в: {docx_blocks}"
        
        last_block = docx_blocks[-1]
        assert "Инциденты информационной безопасности" in last_block, f"Завершающий параграф не найден: {last_block}"
        print(f" [OK] parse_docx успешно извлек {len(docx_blocks)} блоков с сохранением структуры таблиц.")

        # --- 3. Тестирование парсера Excel (.xlsx) ---
        print("\n[3/6] Тестирование parse_xlsx...")
        xlsx_rows = parse_xlsx(str(xlsx_file))
        assert len(xlsx_rows) >= 5, f"Ожидалось >= 5 строк, получено {len(xlsx_rows)}"
        
        # Проверка семантического формата
        for r in xlsx_rows:
            assert r.startswith("Лист: "), f"Строка не начинается с 'Лист: ': {r}"
            assert ": " in r, f"Строка не содержит ключ-значение ': ': {r}"

        # Проверка первого листа (Штатное расписание)
        ivanov_found = any("Иванов Иван Иванович" in r and "Главный архитектор AI" in r for r in xlsx_rows)
        assert ivanov_found, "Строка с Ивановым не найдена в выводе xlsx"

        # Проверка второго листа (Кластер GPU)
        gpu_found = any("Кластер GPU" in r and "Векторизация Nomic" in r for r in xlsx_rows)
        assert gpu_found, "Строка второго листа с кластером GPU не найдена"

        # Проверка отсутствия формул и корректности чистки None
        for r in xlsx_rows:
            assert "=D" not in r, f"Обнаружена невычисленная формула: {r}"
            assert "None" not in r, f"Обнаружена сырая строка 'None': {r}"
        print(f" [OK] parse_xlsx успешно сформировал {len(xlsx_rows)} семантических строк по 2 листам.")

        # --- 4. Тестирование унифицированного диспетчера load_and_chunk_document ---
        print("\n[4/6] Тестирование load_and_chunk_document...")
        docx_chunks = load_and_chunk_document(str(docx_file), chunk_size=300, chunk_overlap=30)
        assert len(docx_chunks) > 0, "load_and_chunk_document вернул пустой список для docx"
        for ch in docx_chunks:
            assert "text" in ch and "metadata" in ch
            assert ch["metadata"]["format"] == "docx"
            assert ch["metadata"]["source"] == docx_file.name
            assert not ch["text"].startswith(NOMIC_DOC_PREFIX), "Сырой текст не должен содержать Nomic-префикс!"

        xlsx_chunks = load_and_chunk_document(str(xlsx_file), chunk_size=200)
        assert len(xlsx_chunks) > 0, "load_and_chunk_document вернул пустой список для xlsx"
        for ch in xlsx_chunks:
            assert "text" in ch and "metadata" in ch
            assert ch["metadata"]["format"] == "xlsx"
            assert ch["metadata"]["sheet"] is not None, "Имя листа Excel отсутствует в метаданных"
            assert not ch["text"].startswith(NOMIC_DOC_PREFIX), "Сырой текст не должен содержать Nomic-префикс!"

        print(f" [OK] load_and_chunk_document сформировал {len(docx_chunks)} docx-чанков и {len(xlsx_chunks)} xlsx-чанков.")

        # --- 5. Проверка Nomic v1.5 протокола форматирования префиксов ---
        print("\n[5/6] Тестирование протокола префиксов Nomic v1.5...")
        sample_raw = "Тестовый фрагмент документа"
        doc_prefixed = format_nomic_document(sample_raw)
        assert doc_prefixed == f"search_document: {sample_raw}", f"Неверный doc префикс: {doc_prefixed}"
        # Идемпотентность
        assert format_nomic_document(doc_prefixed) == doc_prefixed, "Префикс doc задвоился!"

        query_raw = "Где сервер для Nomic?"
        query_prefixed = format_nomic_query(query_raw)
        assert query_prefixed == f"search_query: {query_raw}", f"Неверный query префикс: {query_prefixed}"
        # Идемпотентность и смена типа
        assert format_nomic_query(query_prefixed) == query_prefixed, "Префикс query задвоился!"
        assert format_nomic_query(doc_prefixed) == f"search_query: {sample_raw}", "Перекрестный префикс не скорректирован!"
        print(" [OK] Протокол префиксов Nomic v1.5 строго валидирован.")

        # --- 6. Вычисление эмбеддингов Nomic v1.5 и Semantic Search ---
        print("\n[6/6] Тестирование SentenceTransformer (nomic-embed-text-v1.5) и Dense Retrieval...")
        config = AppConfig(
            embedding_provider="SENTENCE_TRANSFORMERS",
            sentence_transformer_model="nomic-ai/nomic-embed-text-v1.5",
            embedding_device="cpu"
        )
        store = DenseVectorStore(config=config)
        assert store.check_health(), "DenseVectorStore health check не пройден!"

        # Собираем все чанки обоих документов
        all_doc_chunks = docx_chunks + xlsx_chunks
        chunk_texts = [c["text"] for c in all_doc_chunks]
        
        print(f" Расчет эмбеддингов для {len(chunk_texts)} чанков документов (префикс 'search_document: ')...")
        doc_embeddings = store.get_batch_embeddings(chunk_texts, is_query=False)
        assert len(doc_embeddings) == len(chunk_texts)
        
        # Проверка размерности и L2-нормализации
        for vec in doc_embeddings:
            assert vec is not None and len(vec) == 768, f"Размерность вектора {len(vec) if vec else 0} != 768"
            norm = np.linalg.norm(vec)
            assert abs(norm - 1.0) < 1e-4, f"Вектор не нормализован (L2={norm})"

        # Тестовые запросы
        test_queries = [
            {
                "query": "Каковы требования политики паролей к длине символов?",
                "expected_keyword": "16 символов",
                "source_type": "docx"
            },
            {
                "query": "Какие ограничения доступа действуют для роли Администратор?",
                "expected_keyword": "аппаратный 2FA токен",
                "source_type": "docx"
            },
            {
                "query": "Какой сервер в кластере GPU выполняет векторизацию Nomic?",
                "expected_keyword": "SRV-02",
                "source_type": "xlsx"
            },
            {
                "query": "Сколько получает Главный архитектор AI Иванов?",
                "expected_keyword": "Главный архитектор AI",
                "source_type": "xlsx"
            }
        ]

        print("\n Выполнение семантических поисковых запросов:")
        for tq in test_queries:
            q_text = tq["query"]
            q_vec = store.get_embedding(q_text, is_query=True)
            assert q_vec is not None and len(q_vec) == 768
            
            # Подсчет косинусного сходства со всеми чанками
            sims = [cosine_similarity(q_vec, d_vec) for d_vec in doc_embeddings]
            best_idx = int(np.argmax(sims))
            best_sim = sims[best_idx]
            best_chunk = all_doc_chunks[best_idx]
            best_text = best_chunk["text"]

            print(f"\n 🔍 Запрос: '{q_text}'")
            print(f"    Лучшее совпадение: sim={best_sim:.4f} | Источник: {best_chunk['metadata']['source']} ({best_chunk['metadata']['format']})")
            print(f"    Фрагмент: {best_text[:120]}...")

            assert tq["expected_keyword"] in best_text, (
                f"Ошибка семантического поиска! Запрос: '{q_text}'\n"
                f"Ожидалось ключевое слово: '{tq['expected_keyword']}'\n"
                f"Получен фрагмент: '{best_text}'"
            )
            assert best_sim > 0.60, f"Низкая семантическая близость: {best_sim} <= 0.60"

        # --- 7. Тестирование интеграции в DocumentIngestionPipeline ---
        print("\n Проверка сквозной интеграции в DocumentIngestionPipeline.process_directory...")
        ingestion = DocumentIngestionPipeline(chunk_size=500, chunk_overlap=50)
        chunks = ingestion.process_directory(tmp_path, dense_store=store)
        assert len(chunks) > 0, "DocumentIngestionPipeline не вернул чанков"
        docx_in_chunks = any(c.metadata.source == docx_file.name for c in chunks)
        xlsx_in_chunks = any(c.metadata.source == xlsx_file.name for c in chunks)
        assert docx_in_chunks, "DOCX файл не найден среди проиндексированных чанков"
        assert xlsx_in_chunks, "XLSX файл не найден среди проиндексированных чанков"
        
        # Проверка векторов в объектах DocumentChunk
        vectors_computed = all(c.vector is not None and len(c.vector) == 768 for c in chunks)
        assert vectors_computed, "Не у всех DocumentChunk вычислены 768-D векторы"
        print(f" [OK] process_directory успешно обработал папку и создал {len(chunks)} векторизованных чанков.")

    print("\n" + "=" * 70)
    print("✅ [SUCCESS] Все тесты пройдены успешно! Пайплайн DOCX/XLSX + Nomic v1.5 готов к работе.")
    print("=" * 70)


if __name__ == "__main__":
    run_all_tests()
